# Builds the review spreadsheet for the invitationals' own question sets
# (inv_sets.py): one tab per set, with yellow columns to fill in (Keep?, a new
# league, comments), plus the function-picker set for Mobile.
# Usage (from the repo root): python3 tools/gk/make_inv_review.py out.xlsx
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, "tools/gk")
from inv_sets import SETS, picker_rows

ARIAL = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E3D")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
SKILL = {"data": "Data Analysis", "map": "Mapping", "text": "Text Processing", "games": "Game Logic", "math": "Math & Formulas", "time": "Time & Dates", "cards": "Cards & Random"}
LEAGUE = {5: "League 5", 4: "League 4", 3: "League 3", 2: "League 2", 1: "League 1", 0: "Titan"}


def sheet(wb, title, head, rows, widths, inputs):
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows: ws.append(r)
    n = len(rows) + 1
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            # Answers like "=C4 + 1" or "#PYTHON!" are text, not formulas or errors.
            if isinstance(cell.value, str) and cell.value[:1] in "=#":
                cell.data_type = "s"
                cell.quotePrefix = True
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        for c in inputs:
            row[c - 1].fill = INPUT
    keep = get_column_letter(inputs[0])
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{keep}2:{keep}{n}")
    if len(inputs) == 3:
        mv = get_column_letter(inputs[1])
        dv2 = DataValidation(type="list", formula1='"' + ",".join(LEAGUE.values()) + '"', allow_blank=True)
        ws.add_data_validation(dv2)
        dv2.add(f"{mv}2:{mv}{n}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n}"
    ws.row_dimensions[1].height = 30
    return ws, n

wb = Workbook()
g = wb.active
g.title = "How to review"
head = ["Question", "League", "Skill", "Answer", "Wrong 1", "Wrong 2", "Wrong 3", "Source", "Keep?", "Change league to", "Comments (or a better answer)"]
summary = []
TABS = {"pq": "Power Query", "vba": "VBA", "robot": "OA Robot", "lam": "Lambda", "py": "Python", "phone": "Excel on a phone"}
for key, title in TABS.items():
    rows = [[qq, LEAGUE[t], SKILL[sk], a, *w, src] for t, sk, qq, a, w, src in SETS[key]]
    ws, n = sheet(wb, title, head, rows, [46, 10, 15, 34, 26, 26, 26, 24, 9, 16, 40], [9, 10, 11])
    summary.append((title, n, "I", "J", "K"))

prow = [[name, qq, a, *w, desc, LEAGUE[tier]] for name, tier, sk, ti, qq, a, w, desc in picker_rows()]
ws, n = sheet(wb, "Function picker", ["Function", "Question", "Answer", "Wrong 1", "Wrong 2", "Wrong 3", "What it does", "League", "Keep?", "Comments"], prow, [14, 40, 20, 20, 20, 20, 40, 10, 9, 40], [9, 10])
summary.append(("Function picker", n, "I", None, "J"))

g.column_dimensions["A"].width = 30
for c in "BCDE": g.column_dimensions[c].width = 13
def put(r, c, v, bold=False, size=10):
    cell = g.cell(row=r, column=c, value=v)
    cell.font = Font(name=ARIAL, bold=bold, size=size)
    cell.alignment = Alignment(vertical="top")
    return cell
put(1, 1, "CELLOG invitational question sets — review", bold=True, size=14)
lines = [
    "Each invitational's Knowledge check will ask its last (general) question from its own set below, instead of the main pool.",
    "After an invitational's first appearance, its set joins the main general pool (marked with its icon, tracked separately in Career).",
    "Mobile: three 'Excel on a phone' questions per check in place of functions, then one Function picker question in place of the general one.",
    "Only fill in the yellow columns.",
    "Keep?  —  leave blank (or Yes) to keep it; No to drop it.",
    "Change league to  —  only if it's in the wrong league (League 5 = easiest; Titan = hardest). It matters once a set joins the main pool.",
    "Comments  —  a wrong or unclear answer, a better wording, or a question to add.",
    "Source  —  where the fact comes from. 'Good practice' rows are advice rather than a documented step: check those hardest.",
    "Example: Keep? No · Comments: \"The ribbon's at the top on my phone\".",
    "Save it and send the file back in the chat.",
]
for i, t in enumerate(lines):
    put(3 + i, 1, t)
r0 = 4 + len(lines)
put(r0, 1, "Summary", bold=True, size=11)
for c, h in enumerate(["Questions", "Dropped (No)", "Moved", "Comments"], 2):
    put(r0 + 1, c, h, bold=True)
for i, (title, n, kc, mc, cc) in enumerate(summary):
    r = r0 + 2 + i
    put(r, 1, title, bold=True)
    q = f"'{title}'"
    g.cell(row=r, column=2, value=f"=COUNTA({q}!A2:A{n})").font = Font(name=ARIAL)
    g.cell(row=r, column=3, value=f"=COUNTIF({q}!{kc}2:{kc}{n},\"No\")").font = Font(name=ARIAL)
    g.cell(row=r, column=4, value=f"=COUNTA({q}!{mc}2:{mc}{n})" if mc else "-").font = Font(name=ARIAL)
    g.cell(row=r, column=5, value=f"=COUNTA({q}!{cc}2:{cc}{n})").font = Font(name=ARIAL)
r = r0 + 2 + len(summary)
put(r, 1, "Total", bold=True)
for c in (2, 3, 5):
    col = get_column_letter(c)
    g.cell(row=r, column=c, value=f"=SUM({col}{r0 + 2}:{col}{r - 1})").font = Font(name=ARIAL, bold=True)
wb.save(sys.argv[1])
print({k: len(v) for k, v in SETS.items()}, "picker", len(prow))
