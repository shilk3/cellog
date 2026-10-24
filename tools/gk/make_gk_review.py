# Builds a review spreadsheet for the general-knowledge bank (edu-general.js):
# one row per question, with yellow columns to fill in (Keep?, a new league,
# comments). Read back the same way as the function banks' review.
# Usage (from the repo root): python3 tools/gk/make_gk_review.py out.xlsx
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARIAL = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E3D")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
SKILL = {"data": "Data Analysis", "map": "Mapping", "text": "Text Processing", "games": "Game Logic", "math": "Math & Formulas", "time": "Time & Dates", "cards": "Cards & Random"}
LEAGUE = {5: "League 5", 4: "League 4", 3: "League 3", 2: "League 2", 1: "League 1", 0: "Titan"}

js = open("edu-general.js").read()
rows = json.loads("[" + re.search(r"const EDU_GENERAL = \[(.*?)\];", js, re.S).group(1).strip().rstrip(",") + "]")

wb = Workbook()
g = wb.active
g.title = "How to review"
ws = wb.create_sheet("General knowledge")
head = ["Question", "Topic", "League", "Skill", "Answer", "Wrong 1", "Wrong 2", "Wrong 3", "Keep?", "Change league to", "Comments (or a better answer)"]
ws.append(head)
for qid, sk, ti, tier, topic, q, a, w1, w2, w3 in rows:
    ws.append([q, topic, LEAGUE[tier], SKILL[sk], a, w1, w2, w3])
n = len(rows) + 1
for c in range(1, len(head) + 1):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
    cell.fill = HEAD
    cell.alignment = Alignment(vertical="center", wrap_text=True)
for i, w in enumerate([46, 18, 10, 15, 34, 26, 26, 26, 9, 16, 40], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2, max_row=n):
    for cell in row:
        # Answers like "=A2" or "#NUM!" are text, not formulas or errors.
        if isinstance(cell.value, str) and cell.value[:1] in "=#":
            cell.data_type = "s"
            cell.quotePrefix = True
        cell.font = Font(name=ARIAL, size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    for c in (9, 10, 11):
        row[c - 1].fill = INPUT
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"I2:I{n}")
dv2 = DataValidation(type="list", formula1='"' + ",".join(LEAGUE.values()) + '"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"J2:J{n}")
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:K{n}"
ws.row_dimensions[1].height = 30

g.column_dimensions["A"].width = 26
g.column_dimensions["B"].width = 12
def put(r, c, v, bold=False, size=10):
    cell = g.cell(row=r, column=c, value=v)
    cell.font = Font(name=ARIAL, bold=bold, size=size)
    cell.alignment = Alignment(vertical="top")
    return cell
put(1, 1, "CELLOG General Knowledge — review", bold=True, size=14)
lines = [
    "One of these questions ends every Knowledge check. Only fill in the yellow columns.",
    "Keep?  —  leave blank (or Yes) to keep it; No to drop it.",
    "Change league to  —  only if it's in the wrong league (League 5 = easiest; Titan = hardest).",
    "Comments  —  a wrong or unclear answer, a better wording, or a question to add.",
    "Example: Keep? No · Comments: \"Ctrl+Shift+V isn't in my Excel yet\".",
    "Save it and send the file back in the chat.",
]
for i, t in enumerate(lines):
    put(3 + i, 1, t)
put(10, 1, "Summary", bold=True, size=11)
put(11, 1, "Questions", bold=True); g.cell(row=11, column=2, value=f"=COUNTA('General knowledge'!A2:A{n})").font = Font(name=ARIAL)
put(12, 1, "Dropped (No)", bold=True); g.cell(row=12, column=2, value=f"=COUNTIF('General knowledge'!I2:I{n},\"No\")").font = Font(name=ARIAL)
put(13, 1, "Moved", bold=True); g.cell(row=13, column=2, value=f"=COUNTA('General knowledge'!J2:J{n})").font = Font(name=ARIAL)
put(14, 1, "Comments", bold=True); g.cell(row=14, column=2, value=f"=COUNTA('General knowledge'!K2:K{n})").font = Font(name=ARIAL)
wb.save(sys.argv[1])
print(len(rows), "questions")
