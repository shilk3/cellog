# Builds the review spreadsheet for the Power Query groups (groups.py): the
# groups, then every function's "Which one …?" question, with yellow columns
# to fill in. Usage (from the repo root):
#   python3 tools/pq/make_group_review.py pq.json out.xlsx
# where pq.json is EDU_BANKS.pq from edu-banks.js.
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0, "tools/pq")
from groups import G, INTERNAL, CLUE, clue

ARIAL = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E3D")
INPUT = PatternFill("solid", fgColor="FFF2CC")
FLAG = PatternFill("solid", fgColor="FCE4D6")
thin = Side(style="thin", color="D9D9D9")
names = {r[0]: r for r in json.load(open(sys.argv[1]))}
setname = lambda ms: "Completionist" if names[ms[0]][5] else "Esports"

def sheet(wb, title, head, rows, widths, inputs, flagcol=None):
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows: ws.append(r)
    n = len(rows) + 1
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            if isinstance(cell.value, str) and cell.value[:1] in "=#":
                cell.data_type = "s"
                cell.quotePrefix = True
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        for c in inputs:
            row[c - 1].fill = INPUT
        if flagcol and row[flagcol - 1].value:
            row[flagcol - 1].fill = FLAG
    ws.freeze_panes = "B2" if title == "Groups" else "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n}"
    ws.row_dimensions[1].height = 30
    return ws, n

wb = Workbook()
g = wb.active
g.title = "How to review"
grows = [[i + 1, n, setname(ms), len(ms), ", ".join(ms)] for i, (n, ms) in enumerate(G)]
ws, gn = sheet(wb, "Groups", ["#", "Group", "Set", "Functions", "Members", "OK as is?", "Change (merge, split, move a function…)"], grows, [5, 34, 14, 10, 70, 10, 40], [6, 7])
dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"F2:F{gn}")

qrows = []
for n, ms in G:
    for m in ms:
        d = names[m][3]
        same = [x for x in ms if x != m and names[x][3].strip().lower() == d.strip().lower()]
        flag = "Microsoft only says 'internal use only' — suggest leaving it out" if d.strip().rstrip(".") == INTERNAL else f"Same description as {', '.join(same)}" if same else "Hand-written" if m in CLUE else ""
        qrows.append([n, setname(ms), m, clue(m, d), flag, d])
ws, qn = sheet(wb, "Questions", ["Group", "Set", "Function (the answer)", "Question (draft)", "Flag", "Microsoft's description", "Better question", "Move to group", "Comments"], qrows, [30, 13, 30, 60, 26, 50, 40, 22, 30], [7, 8, 9], flagcol=5)

g.column_dimensions["A"].width = 28
g.column_dimensions["B"].width = 14
g.column_dimensions["C"].width = 16
def put(r, c, v, bold=False, size=10):
    cell = g.cell(row=r, column=c, value=v)
    cell.font = Font(name=ARIAL, bold=bold, size=size)
    cell.alignment = Alignment(vertical="top")
    return cell
put(1, 1, "CELLOG Power Query groups — review", bold=True, size=14)
lines = [
    "Power Query is asked in groups of look-alike functions instead of one function at a time.",
    "A group check asks 'Which one …?' for each function in the group; the wrong answers are the others in it.",
    "Get them all right in one go and the group is mastered. Groups are what Career and Completion count.",
    "Each group is all Esports or all Completionist, so a group means the same in both modes.",
    "",
    "Groups tab — is each group a sensible family? Say what to merge, split or move.",
    "Questions tab — one row per function. Questions are drafted from Microsoft's descriptions;",
    "   rewrite any that are clumsy, too easy (the answer's name is in the question) or ambiguous.",
    "   Orange flags need a look first. Only fill in the yellow columns.",
    "Example: Better question: \"Which one returns the last day of the quarter?\"",
    "Save it and send the file back in the chat.",
]
for i, t in enumerate(lines):
    put(3 + i, 1, t)
r0 = 4 + len(lines)
put(r0, 1, "Summary", bold=True, size=11)
for c, h in enumerate(["Esports", "Completionist"], 2):
    put(r0 + 1, c, h, bold=True)
put(r0 + 2, 1, "Groups", bold=True)
put(r0 + 3, 1, "Functions", bold=True)
put(r0 + 4, 1, "Flagged questions", bold=True)
for c, s in ((2, "Esports"), (3, "Completionist")):
    g.cell(row=r0 + 2, column=c, value=f'=COUNTIF(Groups!C2:C{gn},"{s}")').font = Font(name=ARIAL)
    g.cell(row=r0 + 3, column=c, value=f'=SUMIF(Groups!C2:C{gn},"{s}",Groups!D2:D{gn})').font = Font(name=ARIAL)
    g.cell(row=r0 + 4, column=c, value=f'=COUNTIFS(Questions!B2:B{qn},"{s}",Questions!E2:E{qn},"?*")').font = Font(name=ARIAL)
put(r0 + 6, 1, "Your answers so far", bold=True, size=11)
put(r0 + 7, 1, "Groups to change", bold=True)
g.cell(row=r0 + 7, column=2, value=f"=COUNTA(Groups!G2:G{gn})").font = Font(name=ARIAL)
put(r0 + 8, 1, "Questions reworded", bold=True)
g.cell(row=r0 + 8, column=2, value=f"=COUNTA(Questions!G2:G{qn})").font = Font(name=ARIAL)
wb.save(sys.argv[2])
print(len(G), "groups,", len(qrows), "questions,", sum(1 for r in qrows if r[4]), "flagged")
