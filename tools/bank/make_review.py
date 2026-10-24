import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
sys.path.insert(0, "tools/bank")
src = open("tools/bank/build_bank.py").read()
parse_vba_args = {}
exec(src.split("PQLV =")[0], parse_vba_args)  # brings in parse_vba_args
parse_vba_args = parse_vba_args["parse_vba_args"]

ARIAL = "Arial"
HEAD = PatternFill("solid", fgColor="1F4E3D")
INPUT = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
def counts(args):
    mand = sum(1 for a in args if not a.startswith("[") and a not in ("…", "..."))
    opt = sum(1 for a in args if a.startswith("["))
    rep = any(a in ("…", "...") for a in args)
    return mand, (f"{opt} + repeats" if rep else opt)
def argtxt(args):
    return ", ".join("…" if a == "..." else a for a in args) if args else "(none)"

js = open("edu-functions.js").read()
excel = json.loads("[" + re.search(r"const EDU_FUNCS = \[(.*?)\];", js, re.S).group(1).strip().rstrip(",") + "]")
pandas = json.loads("[" + re.search(r"const EDU_PANDAS = \[(.*?)\];", js, re.S).group(1).strip().rstrip(",") + "]")
pq = json.load(open("tools/pq/levelled.json"))
vba = json.load(open("tools/vba/levelled.json"))
rb = json.load(open("tools/robot/robot_clean.json"))
SKILL = {"data": "Data Analysis", "map": "Mapping", "text": "Text Processing", "games": "Game Logic", "math": "Math & Formulas", "time": "Time & Dates", "cards": "Cards & Random"}
LEAGUE = {5: "League 5", 4: "League 4", 3: "League 3", 2: "League 2", 1: "League 1", 0: "Titan"}
PQLV = {1: "Common", 2: "Everyday", 3: "Rare"}
EXCEL_LAMBDA = ["LAMBDA", "LET", "MAP", "REDUCE", "SCAN", "BYROW", "BYCOL", "MAKEARRAY", "ISOMITTED"]

sheets = []
def ex_rows():
    for n, sk, ti, tier, d, a, cat in excel:
        args = a.split("|") if a else []
        m, o = counts(args)
        yield [n, LEAGUE[tier], SKILL[sk], cat, m, o, argtxt(args), d]
sheets.append(dict(name="Excel", head=["Function", "League", "Skill", "Microsoft category", "Mandatory", "Optional", "Arguments", "Description"],
                   rows=list(ex_rows()), change=("Change league to", list(LEAGUE.values())), widths=[18, 10, 16, 18, 10, 12, 44, 60]))
def pq_rows():
    for x in pq:
        args = [("[" + a["n"] + "]") if a["opt"] else a["n"] for a in x["args"]]
        m, o = counts(args)
        yield [x["name"], PQLV[x["level"]], x["cat"].replace("-", " "), m, o, argtxt(args), x["desc"]]
sheets.append(dict(name="Power Query", head=["Function", "Level", "Category", "Mandatory", "Optional", "Arguments", "Description"],
                   rows=list(pq_rows()), change=("Change level to", ["Common", "Everyday", "Rare"]), widths=[34, 10, 16, 10, 12, 44, 60]))
def vba_rows():
    for v in sorted(vba, key=lambda v: v["name"].lower()):
        args = parse_vba_args(v["syntax"])
        m, o = counts(args)
        yield [v["name"], "Common" if v["level"] == 1 else "Other", v["cat"].replace(" functions", ""), m, o, argtxt(args), v["desc"]]
sheets.append(dict(name="VBA", head=["Function", "Level", "Category", "Mandatory", "Optional", "Arguments", "Description"],
                   rows=list(vba_rows()), change=("Change level to", ["Common", "Other"]), widths=[16, 10, 14, 10, 12, 44, 60]))
sheets.append(dict(name="Robot commands", head=["Command", "Tag", "Group", "Description", "From"],
                   rows=[[c["name"], "Robot-only" if c["hard"] else "Easy in Excel", c["group"] or "", c["desc"], c["src"]] for c in rb["cmds"]],
                   change=("Change tag to", ["Robot-only", "Easy in Excel"]), widths=[36, 14, 22, 60, 22],
                   comment_head="Comments (or a description for ones without)"))
def lam_rows():
    for r in excel:
        if r[0] in EXCEL_LAMBDA:
            args = [a if a != "..." else "…" for a in r[5].split("|")] if r[5] else []
            m, o = counts(args)
            yield [r[0], "Excel", m, o, argtxt(args), r[4]]
    for l in rb["lams"]:
        if l["name"] == "ShowHighlights": continue
        args = l["params"] or []
        m, o = counts(args)
        yield [l["name"], l["src"], m, o, argtxt(args), l["desc"]]
sheets.append(dict(name="Lambdas", head=["LAMBDA", "From", "Mandatory", "Optional", "Arguments", "Description"], rows=list(lam_rows()), change=None, widths=[26, 22, 10, 12, 44, 60]))
sheets.append(dict(name="Python", head=["Name", "Kind", "Mandatory", "Optional", "Arguments", "Description"],
                   rows=[[n, "pandas" if n.startswith("pd.") else "DataFrame" if n.startswith("df.") else "numpy" if n.startswith("np.") else "Python in Excel", *counts(a.split("|")), argtxt(a.split("|")), d] for n, d, a in pandas],
                   change=None, widths=[20, 16, 10, 12, 44, 60]))

wb = Workbook()
guide = wb.active
guide.title = "How to review"
for s in sheets:
    ws = wb.create_sheet(s["name"])
    head = s["head"] + ["Keep?"] + ([s["change"][0]] if s["change"] else []) + [s.get("comment_head", "Comments")]
    ws.append(head)
    for r in s["rows"]:
        ws.append(r)
    n = len(s["rows"]) + 1
    base = len(s["head"])
    keep_col = base + 1
    change_col = base + 2 if s["change"] else None
    com_col = base + (3 if s["change"] else 2)
    s.update(keep=keep_col, change_col=change_col, com=com_col, n=n)
    for c in range(1, len(head) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
        cell.fill = HEAD
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = s["widths"] + [9] + ([16] if s["change"] else []) + [40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row:
            cell.font = Font(name=ARIAL, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (base, base - 1) or cell.column == com_col)
            cell.border = Border(bottom=thin)
        for c in [keep_col, change_col, com_col]:
            if c: row[c - 1].fill = INPUT
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, showDropDown=False)
    dv.error, dv.errorTitle = "Pick Yes or No (or leave it blank to keep it).", "Keep?"
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(keep_col)}2:{get_column_letter(keep_col)}{n}")
    if s["change"]:
        dv2 = DataValidation(type="list", formula1='"' + ",".join(s["change"][1]) + '"', allow_blank=True)
        ws.add_data_validation(dv2)
        dv2.add(f"{get_column_letter(change_col)}2:{get_column_letter(change_col)}{n}")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(head))}{n}"
    ws.row_dimensions[1].height = 30

# --- the guide sheet
g = guide
g.column_dimensions["A"].width = 22
for col, w in zip("BCDEF", [12, 14, 16, 14, 50]):
    g.column_dimensions[col].width = w
def put(r, c, v, bold=False, size=10, color=None, fill=None, wrap=False):
    cell = g.cell(row=r, column=c, value=v)
    cell.font = Font(name=ARIAL, bold=bold, size=size, color=color)
    if fill: cell.fill = fill
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    return cell
put(1, 1, "CELLOG Question Bank — review", bold=True, size=14)
put(3, 1, "How to review", bold=True, size=11)
lines = [
    "Each tab is one question bank. Only fill in the yellow columns; leave everything else as it is.",
    "Keep?  —  leave blank to keep a row as it is, or pick No to drop it (Yes is fine too).",
    "Change league / level / tag to  —  pick a new value from the dropdown only if you want to move it.",
    "Comments  —  anything else: a wrong argument, a better description, a question. On Robot commands, you can write a description for ones without.",
    "Filters are on every header row, and the first column stays put while you scroll.",
    "Save it and send the file back in the chat — I'll read every yellow cell and apply them.",
]
for i, t in enumerate(lines):
    put(4 + i, 1, t, wrap=False)
put(11, 1, "Example of a filled-in row (Excel tab)", bold=True, size=11)
ex_head = ["Function", "League", "…", "Keep?", "Change league to", "Comments"]
ex_row = ["ACCRINT", "Titan", "…", "No", "", "Too obscure, even for Titan"]
for i, v in enumerate(ex_head, 1):
    put(12, i, v, bold=True, color="FFFFFF", fill=HEAD)
for i, v in enumerate(ex_row, 1):
    put(13, i, v, fill=INPUT if i >= 4 else None)
put(15, 1, "Summary (updates as you fill it in)", bold=True, size=11)
for i, v in enumerate(["Tab", "Rows", "Dropped (No)", "Moved", "Comments"], 1):
    put(16, i, v, bold=True, color="FFFFFF", fill=HEAD)
for j, s in enumerate(sheets):
    r = 17 + j
    q = f"'{s['name']}'"
    K, C, M = (get_column_letter(s[k]) if s[k] else None for k in ("keep", "change_col", "com"))
    put(r, 1, s["name"])
    put(r, 2, f"=COUNTA({q}!A2:A{s['n']})")
    put(r, 3, f'=COUNTIF({q}!{K}2:{K}{s["n"]},"No")')
    put(r, 4, f"=COUNTA({q}!{C}2:{C}{s['n']})" if C else "—")
    put(r, 5, f"=COUNTA({q}!{M}2:{M}{s['n']})")
tot = 17 + len(sheets)
put(tot, 1, "Total", bold=True)
for c in "BCDE":
    put(tot, "BCDE".index(c) + 2, f"=SUM({c}17:{c}{tot - 1})", bold=True)
put(tot + 2, 1, "Sources: Excel functions — Microsoft's Excel functions (alphabetical) reference; Power Query — Microsoft's Power Query M function reference; VBA — Microsoft's VBA function reference; Robot commands & Lambdas — the OA Robot command collections (people's names removed). Levels, tags and skills are my suggestions for you to check.", wrap=False, size=9, color="595959")
out = sys.argv[1]
wb.save(out)
print({s["name"]: s["n"] - 1 for s in sheets})
